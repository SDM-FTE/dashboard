"""Refresh only JAD rows in the static dashboard JSON from a Graph drive item."""

import json
import math
import os
import sys
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "assets" / "data" / "dashboard-data.json"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"


def required_env(name):
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"Missing required GitHub Actions variable: {name}")
    return value


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def add_numeric(a, b):
    numbers = [value for value in (a, b) if isinstance(value, (int, float)) and not isinstance(value, bool)]
    return sum(numbers) if numbers else None


def download_workbook():
    token = required_env("GRAPH_ACCESS_TOKEN")
    drive_id = quote(required_env("JAD_DRIVE_ID"), safe="")
    item_id = quote(required_env("JAD_ITEM_ID"), safe="")
    url = f"{GRAPH_ROOT}/drives/{drive_id}/items/{item_id}/content"
    request = Request(url, headers={"Authorization": f"Bearer {token}"})
    try:
        with urlopen(request, timeout=90) as response:
            return response.read()
    except HTTPError as error:
        raise RuntimeError(
            f"Microsoft Graph could not download the JAD workbook (HTTP {error.code}). "
            "Check the app permission and the JAD Drive ID and Item ID."
        ) from None
    except URLError as error:
        raise RuntimeError(f"Could not connect to Microsoft Graph: {error.reason}") from None


def extract_jad(workbook_bytes):
    try:
        workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    except Exception as error:
        raise RuntimeError("The downloaded file could not be opened as an Excel workbook.") from error

    try:
        missing_sheets = {"PAK", "sum"} - set(workbook.sheetnames)
        if missing_sheets:
            raise RuntimeError(f"The JAD workbook is missing expected sheet(s): {', '.join(sorted(missing_sheets))}.")

        pak = workbook["PAK"]
        summary_sheet = workbook["sum"]
        summary = {}
        for row in summary_sheet.iter_rows(min_row=5, values_only=True):
            if isinstance(row[0], int) and row[1]:
                summary[str(row[1]).strip()] = {
                    "akPenyetaraan": clean(row[2]),
                    "akSkpPrestasi": clean(row[3]),
                    "totalAkBaru": clean(row[4]),
                }

        values = list(pak.iter_rows(min_row=9, values_only=True))
        records = []
        for index, row in enumerate(values):
            if not isinstance(row[0], int) or not row[1]:
                continue
            name = str(row[1]).strip()
            detail = values[index + 1] if index + 1 < len(values) else ()
            if detail and (detail[0] is not None or detail[1] is not None):
                detail = ()
            years = [add_numeric(row[col], detail[col] if detail else None) for col in range(13, 17)]
            summary_row = summary.get(name, {})
            records.append({
                "no": row[0],
                "name": name,
                "degree": clean(row[2]),
                "publication": clean(row[3]),
                "currentRank": clean(row[4]),
                "rankDate": clean(detail[4]) if detail else None,
                "oldAk": clean(row[5]),
                "proposedRank": clean(row[6]),
                "scienceCluster": clean(row[7]),
                "akEducation": clean(row[8]),
                "akTeaching": clean(row[9]),
                "akResearch": clean(row[10]),
                "akCommunity": clean(row[11]),
                "akSupport": clean(row[12]),
                "skpByYear": years,
                "akPenyetaraan": summary_row.get("akPenyetaraan"),
                "akSkpPrestasi": summary_row.get("akSkpPrestasi"),
                "totalAkBaru": clean(row[17]) or summary_row.get("totalAkBaru"),
                "totalAkLamaBaru": clean(row[18]),
                "hei": [clean(row[col]) for col in range(19, 23)],
                "pakRecommendation": clean(row[23]),
                "pakNote": clean(row[24]),
                "senateRecommendation": clean(row[25]),
                "senateNote": clean(row[26]),
            })

        if not records:
            raise RuntimeError("No lecturer records were found in the JAD workbook; dashboard data was not changed.")
        return records
    finally:
        workbook.close()


def main():
    jad = extract_jad(download_workbook())
    if not DATA_FILE.is_file():
        raise RuntimeError(f"Dashboard JSON was not found at {DATA_FILE.relative_to(ROOT)}.")

    try:
        data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise RuntimeError("The existing dashboard JSON could not be read; it was left unchanged.") from error
    if not isinstance(data, dict) or not isinstance(data.get("meta"), dict) or not isinstance(data.get("bkd"), list):
        raise RuntimeError("The existing dashboard JSON has an unexpected structure; it was left unchanged.")

    data["jad"] = jad
    data["meta"]["jadRows"] = len(jad)
    data["meta"]["jadSource"] = "OneDrive/SharePoint workbook"
    data["meta"]["jadUpdatedAt"] = datetime.now(timezone.utc).replace(microsecond=0).isoformat()

    temporary = DATA_FILE.with_suffix(".json.tmp")
    try:
        temporary.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        temporary.replace(DATA_FILE)
    finally:
        if temporary.exists():
            temporary.unlink()

    print(f"Updated JAD rows: {len(jad)}. Preserved BKD rows: {len(data['bkd'])}.")


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as error:
        print(f"JAD refresh failed: {error}", file=sys.stderr)
        raise SystemExit(1)
